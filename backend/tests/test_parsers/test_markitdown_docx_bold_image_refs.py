"""DOCX 加粗/斜体运行内图片引用保留的回归测试。

生产事故（2026-10-19，review 任务 110fa702，用户 17320163035）：投标文件把
ISO27001/ITSS 等证书扫描件放在加粗格式的段落里，mammoth 输出
``<strong><img/></strong>``，而 markdownify 的 ``keep_inline_images_in`` 只检查
img 的**直接父标签**且当时不含 ``strong`` —— ``![alt](path)`` 引用被静默丢弃、
正文只剩 alt 文字，审查 AI 据此误报「证书扫描件缺失，仅文件名占位文字」，
错估损失 4 分（图片实际已落盘为 image_129/135/139/141，仅 markdown 丢引用）。

修复分两层：
1. ``_KEEP_INLINE_IMAGES_IN`` 覆盖 mammoth 会包行内图片的全部常见标签；
2. ``_repair_dropped_image_refs`` 兜底：落盘未引用的图片按唯一 alt 原位补回
   引用，否则在文末补录——任何未来的遗漏都不再静默发生。
"""

import sys
from pathlib import Path

project_root = Path(__file__).resolve().parents[3]
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

from backend.parsers.markitdown_converter import (  # noqa: E402
    DirectFileImageHandler,
    ImageInfo,
    MarkitdownConverter,
    _repair_dropped_image_refs,
)

ALT = "郑州迪维勒普科技有限公司I1(1)_01"


def _make_png(path: Path, color: str) -> Path:
    from PIL import Image

    Image.new("RGB", (32, 16), color).save(path)
    return path


def _add_run_image(container_paragraph, png_path: Path, *, bold=False, italic=False, descr=""):
    run = container_paragraph.add_run()
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    run.add_picture(str(png_path))
    if descr:
        doc_pr = run._r.xpath(".//wp:docPr")[0]
        doc_pr.set("descr", descr)
    return run


def _convert(tmp_path: Path, build_doc) -> tuple:
    from docx import Document

    docx_path = tmp_path / "fixture.docx"
    document = Document()
    build_doc(document, tmp_path)
    document.save(docx_path)

    images_dir = tmp_path / "fixture_images"
    result = MarkitdownConverter().convert(docx_path, images_dir=images_dir)
    return result, images_dir


def test_bold_run_image_in_table_keeps_ref_in_place(tmp_path: Path):
    """事故原样：加粗段落 + 表格单元格内的证书扫描件，引用必须保留。"""

    def build(document, tmp):
        png = _make_png(tmp / "cert.png", "red")
        table = document.add_table(rows=1, cols=1)
        _add_run_image(table.rows[0].cells[0].paragraphs[0], png, bold=True, descr=ALT)
        document.add_paragraph("证书查询截图：")

    result, images_dir = _convert(tmp_path, build)

    assert len(result.images) == 1
    assert (images_dir / result.images[0].filename).is_file()
    expected_ref = f"![{ALT}]({images_dir.name}/{result.images[0].filename})"
    assert expected_ref in result.markdown_content
    # 引用未被降级成裸 alt 文字（事故表象）
    assert f"**{ALT}**" not in result.markdown_content


def test_italic_run_image_keeps_ref(tmp_path: Path):
    def build(document, tmp):
        png = _make_png(tmp / "cert.png", "green")
        _add_run_image(document.add_paragraph(), png, italic=True, descr="ITSS三级-副本")

    result, images_dir = _convert(tmp_path, build)

    assert len(result.images) == 1
    assert f"![ITSS三级-副本]({images_dir.name}/{result.images[0].filename})" in (
        result.markdown_content
    )


def test_bold_run_image_without_descr_keeps_ref(tmp_path: Path):
    """空 alt 的加粗图片：白名单修复后引用同样直接保留（无需兜底）。"""

    def build(document, tmp):
        png = _make_png(tmp / "cert.png", "blue")
        _add_run_image(document.add_paragraph(), png, bold=True)

    result, images_dir = _convert(tmp_path, build)

    assert len(result.images) == 1
    assert f"]({images_dir.name}/{result.images[0].filename})" in result.markdown_content
    assert "解析兜底" not in result.markdown_content


def test_repair_splices_unique_escaped_alt_in_place(tmp_path: Path):
    """markdownify 转义过的 alt（\\* \\_）也要能唯一匹配并原位补回。"""
    handler = DirectFileImageHandler(tmp_path, "doc_images")
    handler.images = [ImageInfo(filename="image_1.png", data=b"x")]
    handler.alt_by_filename = {"image_1.png": "证书*A_1"}

    markdown = "| 标题 |\n| --- |\n| **证书\\*A\\_1** |\n"
    out = _repair_dropped_image_refs(markdown, handler, "doc_images")

    assert "**![证书*A_1](doc_images/image_1.png)**" in out
    assert "解析兜底" not in out


def test_repair_appends_when_alt_ambiguous(tmp_path: Path):
    """alt 在正文出现多次时不做原位替换，统一文末补录。"""
    handler = DirectFileImageHandler(tmp_path, "doc_images")
    handler.images = [ImageInfo(filename="image_1.png", data=b"x")]
    handler.alt_by_filename = {"image_1.png": "同一文字"}

    markdown = "前文提到同一文字，后文又提到同一文字。\n"
    out = _repair_dropped_image_refs(markdown, handler, "doc_images")

    assert "解析兜底" in out
    assert "![同一文字](doc_images/image_1.png)" in out
    # 原文两处保持原样，未被替换
    assert out.count("提到同一文字") == 2


def test_repair_noop_when_everything_referenced(tmp_path: Path):
    handler = DirectFileImageHandler(tmp_path, "doc_images")
    handler.images = [ImageInfo(filename="image_1.png", data=b"x")]
    handler.alt_by_filename = {"image_1.png": "已引用"}

    markdown = "正文 ![已引用](doc_images/image_1.png) 结束\n"
    out = _repair_dropped_image_refs(markdown, handler, "doc_images")

    assert out == markdown
