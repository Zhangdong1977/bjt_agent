"""DOCX 盲区结构探测/恢复 + alt 特殊字符转义的回归测试。

背景（2026-10-21 评估，doc/workspace/21 号 §八）：mammoth 只读 document.xml
正文，对 SmartArt/chart 的 graphicData 零告警丢弃、mc:Choice-only 形状整块
消失、页眉/页脚 part 完全不读、m:oMath 丢弃仅留日志 warning——这些结构里的
审查材料"无声消失"，agent 只能误判「未提供」。生产存量实测频率（1919 份）：
SmartArt 22、页眉图 293、公式 15、Choice-only 2、chart 0、alt 含方括号 42。

修复分三层：
1. ``_scan_docx_blind_spots``：字节级探测 + md 文末「解析覆盖提示」标注，
   SmartArt/公式文字与页眉图片做有上限的廉价恢复；
2. ``_escape_alt_text``：alt 按 CommonMark 反斜杠转义（markdownify 对 alt
   零转义直传，裸 ``]`` 会让 IMAGE_REF_RE 消费方截断失配）；
3. ``parse_quality``：结构化质量信号（document_parser 持久化到
   documents.source_metadata，不进 md）。
"""

import sys
import zipfile
from pathlib import Path

project_root = Path(__file__).resolve().parents[3]
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

from backend.parsers.markitdown_converter import (  # noqa: E402
    DirectFileImageHandler,
    ImageInfo,
    MarkitdownConverter,
    _escape_alt_text,
    _repair_dropped_image_refs,
)
from backend.utils.markdown_images import IMAGE_REF_RE  # noqa: E402

PNG_BYTES = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c63f8cfc0f01f0005050201edb56eab"
    "0000000049454e44ae426082"
)

NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'mc:Ignorable="w14"'
)

CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Default Extension="png" ContentType="image/png"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
{extra}</Types>"""

RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

DOC_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rIdImg1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>
{extra}</Relationships>"""


def _inline_img(pr_id: int, descr: str, rid: str = "rIdImg1") -> str:
    return (
        f'<w:drawing><wp:inline><wp:extent cx="9525" cy="9525"/>'
        f'<wp:docPr id="{pr_id}" name="Pic{pr_id}" descr="{descr}"/>'
        f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{pr_id}" name="Pic{pr_id}"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="9525" cy="9525"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
        f'</a:graphicData></a:graphic></wp:inline></w:drawing>'
    )


def _para(inner: str) -> str:
    return f"<w:p><w:r>{inner}</w:r></w:p>"


def _build_docx(
    path: Path,
    body: str,
    doc_rels_extra: str = "",
    content_extra: str = "",
    extra_parts: dict | None = None,
) -> Path:
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f"<w:document {NS}><w:body>{body}"
        '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>'
        "</w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", CONTENT_TYPES.format(extra=content_extra))
        zf.writestr("_rels/.rels", RELS)
        zf.writestr("word/document.xml", document)
        zf.writestr("word/_rels/document.xml.rels", DOC_RELS.format(extra=doc_rels_extra))
        zf.writestr("word/media/image1.png", PNG_BYTES)
        for part_name, data in (extra_parts or {}).items():
            zf.writestr(part_name, data)
    return path


def _convert(tmp_path: Path, name: str, body: str, **kwargs):
    docx_path = _build_docx(tmp_path / f"{name}.docx", body, **kwargs)
    images_dir = tmp_path / f"{name}_images"
    result = MarkitdownConverter().convert(docx_path, images_dir=images_dir)
    return result, images_dir


# ---- alt 特殊字符（潜伏缺陷：IMAGE_REF_RE 消费方截断失配）----


def test_alt_with_brackets_is_escaped_and_survives_consumer(tmp_path: Path):
    """alt 含 ``[``/``]``/``()``：转义后完整引用，IMAGE_REF_RE 正确取到 src。"""
    result, images_dir = _convert(
        tmp_path,
        "alt_bracket",
        _para('<w:t>证明材料：</w:t>' + _inline_img(1, "证书[批注](第3版)")),
    )

    assert len(result.images) == 1
    md = result.markdown_content
    match = IMAGE_REF_RE.search(md)
    assert match is not None, md
    assert match.group(2) == f"{images_dir.name}/{result.images[0].filename}"
    # 转义形式的引用在 md 中，裸 alt 不再出现
    assert "证书\\[批注\\](第3版)" in md


def test_repair_section_refs_also_escape_alt(tmp_path: Path):
    """兜底补录的引用同样要转义 alt（repair 的两个产出点）。"""
    handler = DirectFileImageHandler(tmp_path, "doc_images")
    handler.images = [ImageInfo(filename="image_1.png", data=b"x")]
    handler.alt_by_filename = {"image_1.png": "证书[批注]"}

    out = _repair_dropped_image_refs("正文没有该图的占位文字。\n", handler, "doc_images")

    assert "![证书\\[批注\\]](doc_images/image_1.png)" in out
    assert IMAGE_REF_RE.search(out).group(2) == "doc_images/image_1.png"


def test_escape_alt_handles_backslash_and_newline():
    assert _escape_alt_text("a[b]\\c") == "a\\[b\\]\\\\c"
    assert _escape_alt_text("第一行\r\n第二行") == "第一行 第二行"


# ---- 盲区结构：探测 + 标注 + 廉价恢复 ----


def test_smartart_text_recovered_with_note(tmp_path: Path):
    """SmartArt：图内文字补录 + 「解析覆盖提示」标注（不得判「未提供」）。"""
    diagram_data = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<dgm:dataModel xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        "<dgm:ptLst>"
        '<dgm:pt><dgm:t><a:p><a:r><a:t>总经理</a:t></a:r></a:p></dgm:t></dgm:pt>'
        '<dgm:pt><dgm:t><a:p><a:r><a:t>技术部</a:t></a:r></a:p></dgm:t></dgm:pt>'
        "</dgm:ptLst></dgm:dataModel>"
    )
    body = _para(
        '<w:t>组织架构：</w:t>'
        '<w:drawing><wp:inline><wp:docPr id="2" name="SA"/>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
        "</a:graphicData></a:graphic></wp:inline></w:drawing>"
    )
    result, _ = _convert(
        tmp_path, "smartart", body, extra_parts={"word/diagrams/data1.xml": diagram_data}
    )

    md = result.markdown_content
    assert "解析覆盖提示" in md
    assert "SmartArt 图示 1 处" in md
    assert "总经理、技术部" in md
    assert result.parse_quality["blind_spots"] == {"smartart": 1}


def test_header_image_materialized_with_note(tmp_path: Path):
    """页眉图片：落盘 + 引用 + 标注；页眉里的超链接关系不得被当图片提取。"""
    header = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f"<w:hdr {NS}><w:p><w:r><w:t>公司抬头</w:t></w:r>"
        f"<w:r>{_inline_img(9, '页眉证书图')}</w:r></w:p></w:hdr>"
    )
    header_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdImg1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>'
        '<Relationship Id="rIdSite" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com"/>'
        "</Relationships>"
    )
    result, images_dir = _convert(
        tmp_path,
        "header_img",
        _para("<w:t>正文第一段</w:t>"),
        doc_rels_extra=(
            '<Relationship Id="rIdHdr" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.xml"/>'
        ),
        content_extra=(
            '<Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>'
        ),
        extra_parts={
            "word/header1.xml": header,
            "word/_rels/header1.xml.rels": header_rels,
        },
    )

    md = result.markdown_content
    assert "页眉/页脚图片 1 张" in md
    header_files = [p.name for p in images_dir.glob("headerimg_*")]
    assert header_files == [result.images[0].filename]
    assert (images_dir / result.images[0].filename).read_bytes() == PNG_BYTES
    assert f"![页眉/页脚图片 {result.images[0].filename}]({images_dir.name}/{result.images[0].filename})" in md
    assert result.parse_quality["blind_spots"] == {"header_footer_images": 1}


def test_omml_math_text_recovered(tmp_path: Path):
    body = (
        "<w:p><w:r><w:t>公式：</w:t></w:r>"
        "<m:oMath xmlns:m=\"http://schemas.openxmlformats.org/officeDocument/2006/math\">"
        "<m:r><m:t>E=mc</m:t></m:r><m:r><m:t>2</m:t></m:r></m:oMath></w:p>"
    )
    result, _ = _convert(tmp_path, "omml", body)

    md = result.markdown_content
    assert "数学公式 1 处" in md
    assert "E=mc" in md
    assert result.parse_quality["blind_spots"] == {"math": 1}


def test_choice_only_shape_noted(tmp_path: Path):
    """无 VML Fallback 的 mc:AlternateContent：整块内容消失但必须留下标注。"""
    body = _para(
        '<mc:AlternateContent>'
        '<mc:Choice Requires="wps"><w:t>Choice侧内容</w:t></mc:Choice>'
        "</mc:AlternateContent>"
    )
    result, _ = _convert(tmp_path, "choice_only", body)

    md = result.markdown_content
    assert "仅现代格式(Choice)的形状/文本框 1 处" in md
    assert result.parse_quality["blind_spots"] == {"choice_only_shapes": 1}


def test_chart_object_noted(tmp_path: Path):
    body = _para(
        '<w:t>业绩图：</w:t>'
        '<w:drawing><wp:inline><wp:docPr id="3" name="Chart"/>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart r:id="rIdChart"/></a:graphicData></a:graphic></wp:inline></w:drawing>'
    )
    result, _ = _convert(tmp_path, "chart", body)

    md = result.markdown_content
    assert "图表对象 1 处" in md
    assert result.parse_quality["blind_spots"] == {"chart": 1}


def test_chart_namespace_declaration_alone_is_not_a_blind_spot(tmp_path: Path):
    """根元素声明了 chart/diagram 命名空间但没有对应 graphicData：不算盲区。"""
    result, _ = _convert(tmp_path, "ns_only", _para("<w:t>只有命名空间声明</w:t>"))

    assert "解析覆盖提示" not in result.markdown_content
    assert result.parse_quality["blind_spots"] == {}


# ---- parse_quality 基线 ----


def test_clean_docx_has_no_notes_but_quality_populated(tmp_path: Path):
    from docx import Document

    docx_path = tmp_path / "clean.docx"
    document = Document()
    run = document.add_paragraph().add_run("普通段落")
    run.bold = True
    document.add_paragraph("纯文本第二段")
    document.save(docx_path)

    images_dir = tmp_path / "clean_images"
    result = MarkitdownConverter().convert(docx_path, images_dir=images_dir)

    assert "解析覆盖提示" not in result.markdown_content
    assert "解析兜底" not in result.markdown_content
    quality = result.parse_quality
    assert quality is not None
    assert quality["images_total"] == 0
    assert quality["blind_spots"] == {}
    assert "dropped_refs_repaired" in quality
    assert "mammoth_warnings" in quality
